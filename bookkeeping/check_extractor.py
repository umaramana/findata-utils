"""
Check Image OCR Extractor — core logic (no Streamlit).

Pipeline: load pages -> find check boxes (OpenCV) -> Surya OCR (local, CPU)
-> field extraction (position + regex) -> review rows -> Excel/CSV export.

Statement pages (a grid of checks, each with a printed "Check# / date / $amount"
line under it) get ONE OCR pass per page; check no., date and amount come from
the printed line, and only the payee is read from handwriting.

Field extraction works on plain OCR lines (text, bbox, confidence), so it is
engine-independent and testable without Surya installed.
Spec: bookkeeping/check_ocr_spec.md
"""
import io
import re
from dataclasses import dataclass, field

import numpy as np
from PIL import Image, ImageOps

SUPPORTED_TYPES = ["png", "jpg", "jpeg", "pdf"]
PDF_DPI = 200
CONF_LEVELS = ["HIGH", "MEDIUM", "LOW"]


@dataclass
class OcrLine:
    text: str
    bbox: tuple          # (x0, y0, x1, y1) in pixels of the check image
    confidence: float = 0.0


@dataclass
class CheckImage:
    source: str          # original file name
    page: int            # 1-based page number within the file
    index: int           # 1-based check number within the page
    image: Image.Image = field(repr=False)


# ── Loading ───────────────────────────────────────────────────────────────────

def load_pages(file_name, data):
    """Return [(page_no, PIL.Image RGB)] for an image or PDF given as bytes."""
    if file_name.lower().endswith(".pdf"):
        import pymupdf as fitz  # pip-only, no poppler needed
        pages = []
        with fitz.open(stream=data, filetype="pdf") as doc:
            zoom = PDF_DPI / 72
            for i, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                pages.append((i, Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
        return pages
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img).convert("RGB")  # phone photos carry rotation in EXIF
    return [(1, img)]


# ── Segmentation ──────────────────────────────────────────────────────────────

MIN_ASPECT, MAX_ASPECT = 1.8, 3.4   # spec: ~2:1 to 3:1, with slack for skew/cropping
MIN_AREA_FRAC = 0.008  # a 24-check statement grid is ~2% per check; 17-up measured at ~3%
SIZE_BAND = (0.5, 2.0)  # keep boxes within this factor of the median check-shaped box


def segment_checks(page_img):
    """Split a page into check images; the whole page when no check box is found."""
    boxes = find_check_boxes(page_img)
    if not boxes:
        return [page_img]
    return [page_img.crop((x, y, x + w, y + h)) for x, y, w, h in boxes]


def find_check_boxes(page_img):
    """Check rectangles (x, y, w, h), row by row (top-to-bottom, then left-to-right)."""
    import cv2
    rgb = np.array(page_img)
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    blur = cv2.GaussianBlur(gray, (5, 5), 0)
    edges = cv2.Canny(blur, 30, 120)
    edges = cv2.dilate(edges, np.ones((5, 5), np.uint8), iterations=2)
    contours, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    page_h, page_w = gray.shape
    page_area = page_w * page_h
    boxes = []
    for c in contours:
        x, y, w, h = cv2.boundingRect(c)
        if h == 0 or w * h < MIN_AREA_FRAC * page_area or w * h > 0.98 * page_area:
            continue
        if MIN_ASPECT <= w / h <= MAX_ASPECT:
            boxes.append((x, y, w, h))

    boxes = _same_size_as_most(_drop_nested(boxes))
    boxes.sort(key=lambda b: (b[1] // max(b[3] // 2, 1), b[0]))  # row bands tolerate a few px of skew
    return boxes


def _same_size_as_most(boxes):
    """Checks on one page share a size; drop odd ones out (logos, header bars)."""
    if len(boxes) < 3:
        return boxes
    med = float(np.median([w * h for _, _, w, h in boxes]))
    lo, hi = SIZE_BAND
    return [b for b in boxes if lo * med <= b[2] * b[3] <= hi * med]


def _drop_nested(boxes):
    """Remove boxes mostly inside a larger box (e.g. the amount box inside a check)."""
    kept = []
    for b in sorted(boxes, key=lambda b: b[2] * b[3], reverse=True):
        if not any(_overlap_frac(b, k) > 0.8 for k in kept):
            kept.append(b)
    return kept


def _overlap_frac(inner, outer):
    ix0, iy0 = max(inner[0], outer[0]), max(inner[1], outer[1])
    ix1 = min(inner[0] + inner[2], outer[0] + outer[2])
    iy1 = min(inner[1] + inner[3], outer[1] + outer[3])
    if ix1 <= ix0 or iy1 <= iy0:
        return 0.0
    return (ix1 - ix0) * (iy1 - iy0) / float(inner[2] * inner[3])


# ── OCR engine (Surya, lazy-loaded) ───────────────────────────────────────────

class OcrEngineError(RuntimeError):
    pass


class SuryaEngine:
    """Wraps Surya 0.17.x (pure torch). Models load once; first run downloads them."""

    def __init__(self):
        try:
            from surya.detection import DetectionPredictor
            from surya.foundation import FoundationPredictor
            from surya.recognition import RecognitionPredictor
        except ImportError as e:
            raise OcrEngineError(
                "Surya OCR is not installed. Run: pip install -r bookkeeping/requirements.txt"
            ) from e
        try:
            self._det = DetectionPredictor()
            self._rec = RecognitionPredictor(FoundationPredictor())
        except OSError as e:  # disk full / no network on first-run model download
            raise OcrEngineError(
                "Could not download or read the OCR models (first run needs internet and ~2 GB "
                f"free disk). ({type(e).__name__}: {e})"
            ) from e
        except (MemoryError, RuntimeError) as e:
            raise OcrEngineError(
                "Could not load the OCR models — likely not enough free RAM. "
                f"Close other apps and retry. ({type(e).__name__}: {e})"
            ) from e

    def ocr(self, img):
        """Return [OcrLine] for one check image."""
        result = self._rec([img], det_predictor=self._det, math_mode=False)[0]
        return [
            OcrLine(text=_strip_tags(ln.text), bbox=tuple(ln.bbox), confidence=float(ln.confidence or 0))
            for ln in result.text_lines
            if ln.text and ln.text.strip()
        ]


def _strip_tags(text):
    # Surya may wrap styled text in <b>/<i> etc.
    return re.sub(r"</?[a-z]+>", "", text).strip()


# ── Field extraction ──────────────────────────────────────────────────────────

DATE_RE = re.compile(
    r"\b(\d{1,2}\s*[/\-.]\s*\d{1,2}\s*[/\-.]\s*\d{2,4}"
    r"|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{1,2},?\s+\d{2,4})\b",
    re.I,
)
AMOUNT_RE = re.compile(r"\$?\s*(\d{1,3}(?:[,\s]\d{3})+|\d+)(?:\s*[.,]\s*(\d{2}))?(?!\d)")
CHECK_NO_RE = re.compile(r"^\s*(?:no\.?|#)?\s*(\d{3,6})\s*$", re.I)
# greedy over each optional word, so a split label ("PAY TO THE" / "ORDER OF") leaves no "THE" behind
PAYEE_LABEL_RE = re.compile(r"pay\s*to(?:\s*the)?(?:\s*order)?(?:\s*of)?\b", re.I)
PAYEE_PREFIX_RE = re.compile(
    r"^\s*(?:pay\s*to(?:\s*the)?(?:\s*order\s*of)?|(?:the\s*)?order\s*of)\b[\s:.\-]*", re.I)
CAPTION_NO_RE = re.compile(r"(?:check|chk|ck)?\s*(?:#|no\.?|number)\s*:?\s*(\d{1,8})\b", re.I)
CAPTION_AMT_RE = re.compile(r"\$\s*(-?\d[\d,]*\.\d{2})\b")
LABEL_FRAGMENT_RE = re.compile(r"^\s*(?:the\s*)?(?:order\s*)?(?:of)?\s*[:.]?\s*$", re.I)
MEMO_LABEL_RE = re.compile(r"^\s*(memo|for|re)\b[:.]?", re.I)
DATE_LABEL_RE = re.compile(r"^\s*date\b[:.]?", re.I)
DOLLARS_WORD_RE = re.compile(r"\bdollars\b", re.I)
MICR_RE = re.compile(r"[⑆⑈⑇⑉]|^[\d\s:;|]{12,}$")


def extract_fields(lines, width, height):
    """Map OCR lines of one check to fields.

    Returns dict: check_no, date, amount (float|None), payee, purpose,
    ocr_conf (mean line confidence), raw_text.
    """
    nl = [_norm(ln, width, height) for ln in lines]
    nl.sort(key=lambda l: (round(l["cy"], 2), l["x0"]))
    used, src = set(), {}  # src: field -> OCR confidence of the line the value came from

    amount = _find_amount(nl, used, src)  # first: the amount box may share the payee line
    payee = (_payee_from_amount_row(nl, used, src)
             or _value_after_label(nl, PAYEE_LABEL_RE, used, src, "payee"))
    purpose = _value_after_label(nl, MEMO_LABEL_RE, used, src, "purpose", region=lambda l: l["cy"] > 0.55)
    date = _find_date(nl, used, src)
    check_no = _find_check_no(nl, used)

    confs = [l["conf"] for l in nl]
    return {
        "check_no": check_no,
        "date": date,
        "amount": amount,
        "payee": payee,
        "purpose": purpose,
        "ocr_conf": round(sum(confs) / len(confs), 3) if confs else 0.0,
        # weakest of the values that matter most; printed labels would inflate a plain mean
        "value_conf": round(min(src.get("amount", 0), src.get("payee", 0)), 3),
        "amount_conf": round(src.get("amount", 0), 3),
        "raw_text": "\n".join(l["text"] for l in nl),
    }


def _norm(ln, w, h):
    x0, y0, x1, y1 = ln.bbox
    return {
        "text": ln.text.strip(), "conf": ln.confidence,
        "x0": x0 / w, "x1": x1 / w, "y0": y0 / h, "y1": y1 / h,
        "cy": (y0 + y1) / 2 / h, "id": id(ln),
    }


def _neighbours_right(nl, anchor, used, max_dy=0.08):
    """Unused lines on roughly the same row as anchor, to its right, nearest first."""
    cands = [
        l for l in nl
        if l["id"] not in used and l is not anchor
        and abs(l["cy"] - anchor["cy"]) <= max_dy and l["x0"] >= anchor["x1"] - 0.02
    ]
    return sorted(cands, key=lambda l: l["x0"])


def _value_after_label(nl, label_re, used, src, name, region=None):
    """Text after a printed label on the same line, else the nearest line to its right."""
    for l in nl:
        if region and not region(l):
            continue
        m = label_re.search(l["text"])
        if not m:
            continue
        used.add(l["id"])
        rest = _clean_value(l["text"][m.end():])
        if LABEL_FRAGMENT_RE.match(rest):
            rest = ""
        if rest:
            src[name] = l["conf"]
            return rest
        for cand in _neighbours_right(nl, l, used):
            if AMOUNT_RE.fullmatch(cand["text"].replace("$", "").strip()):
                continue  # the amount box sits on the payee row
            if LABEL_FRAGMENT_RE.match(cand["text"]):
                continue  # second half of a split label
            used.add(cand["id"])
            src[name] = cand["conf"]
            return _clean_value(cand["text"])
        return ""
    return ""


def _payee_from_amount_row(nl, used, src):
    """Payee = everything left of the '$' amount on its row, label words stripped.

    OCR splits this row unpredictably ("PAY TO THE" + name, "ORDER OF" + name,
    a two-line label, or label/name/amount on one line), so join the whole row.
    """
    amt = src.get("amount_line")
    if amt is None:
        return ""
    tol = max(0.6 * (amt["y1"] - amt["y0"]), 0.03)
    row = [l for l in nl
           if l is not amt and l["id"] not in used and abs(l["cy"] - amt["cy"]) <= tol
           and l["x1"] <= amt["x0"] + 0.02 and not DATE_RE.search(l["text"])]
    row.sort(key=lambda l: (round(l["x0"] / 0.02), l["cy"]))  # stacked label halves: top first
    parts = [l["text"] for l in row]
    pre = amt["text"].split("$")[0] if "$" in amt["text"] else ""
    if re.search(r"[A-Za-z]{2}", pre):
        parts.append(pre)  # "Pay to ... John Smith $ 85.50"; skips OCR junk like "1$7,645.49"
    text = " ".join(parts)
    for _ in range(2):  # "PAY TO THE" then a separate "ORDER OF"
        text = PAYEE_PREFIX_RE.sub("", text)
    text = _clean_value(text)
    if not text:
        return ""
    used.update(l["id"] for l in row)
    src["payee"] = min([l["conf"] for l in row] or [amt["conf"]])
    return text


def _clean_value(text):
    text = re.sub(r"^[\s:.\-_]+|[\s_]+$", "", text)
    text = re.sub(r"\s*\$\s*[\d,]+(?:\.\d{2})?\s*$", "", text)  # amount box merged into payee line
    text = re.sub(r"\s*\$\s*$", "", text)  # lone "$" of the amount box
    return text.strip()


def _find_date(nl, used, src):
    for l in nl:  # labelled date first: "DATE 3/15/24" or "DATE" + handwriting to the right
        if DATE_LABEL_RE.search(l["text"]):
            m = DATE_RE.search(l["text"])
            if m:
                used.add(l["id"])
                src["date"] = l["conf"]
                return m.group(0)
            for cand in _neighbours_right(nl, l, used):
                m = DATE_RE.search(cand["text"])
                if m:
                    used.update({l["id"], cand["id"]})
                    src["date"] = cand["conf"]
                    return m.group(0)
    for l in nl:  # any date in the top half
        if l["id"] in used or l["cy"] > 0.5:
            continue
        m = DATE_RE.search(l["text"])
        if m:
            used.add(l["id"])
            src["date"] = l["conf"]
            return m.group(0)
    return ""


def _parse_amount(text):
    if "$" in text:
        text = text[text.index("$"):]  # "Unit 4 Rentals $ 85.50" -> the $ figure, not the 4
    m = AMOUNT_RE.search(text)
    if not m:
        return None
    whole = re.sub(r"[,\s]", "", m.group(1))
    cents = m.group(2) or "00"
    try:
        return float(f"{whole}.{cents}")
    except ValueError:
        return None


def _find_amount(nl, used, src):
    """Numeric amount: prefer '$'-marked text, right half, upper 70% of the check."""
    best, best_score = None, -1
    for l in nl:
        t = l["text"]
        if l["id"] in used or l["cy"] > 0.7 or MICR_RE.search(t) or DATE_RE.search(t):
            continue
        if DOLLARS_WORD_RE.search(t) and "$" not in t:
            continue  # the written-out legal line
        val = _parse_amount(t)
        if val is None or val == 0:
            continue
        score = 0
        score += 3 if "$" in t else 0
        score += 2 if l["x0"] > 0.55 else 0
        score += 1 if re.search(r"[.,]\s*\d{2}\b", t) else 0
        score += 1 if 0.2 < l["cy"] < 0.6 else 0
        if CHECK_NO_RE.match(t) and l["cy"] < 0.2:
            score -= 3  # looks like the check number
        if score > best_score:
            best, best_score = (l, val), score
    if best and best_score >= 2:
        used.add(best[0]["id"])
        src["amount"] = best[0]["conf"]
        src["amount_line"] = best[0]
        return best[1]
    return None


def _find_check_no(nl, used):
    tops = [l for l in nl if l["id"] not in used and l["cy"] < 0.3 and l["x0"] > 0.6]
    for l in sorted(tops, key=lambda l: (l["cy"], -l["x0"])):
        m = CHECK_NO_RE.match(l["text"])
        if m:
            used.add(l["id"])
            return m.group(1)
    return ""


# ── Statement pages: one OCR pass, printed caption under each check ─────────

def extract_page(lines, boxes):
    """Fields for each check box on a page OCR'd as a whole.

    lines: OcrLines in page pixels. Returns [(fields, view_box)], where view_box
    is the check plus its printed caption band (for the review image).
    """
    out = []
    for i, (x, y, w, h) in enumerate(boxes):
        band = _caption_band(boxes, i)
        inside, caption = [], []
        for ln in lines:
            cx, cy = (ln.bbox[0] + ln.bbox[2]) / 2, (ln.bbox[1] + ln.bbox[3]) / 2
            if not x <= cx <= x + w:
                continue
            if y <= cy <= y + h:
                inside.append(OcrLine(ln.text, (ln.bbox[0] - x, ln.bbox[1] - y,
                                                ln.bbox[2] - x, ln.bbox[3] - y), ln.confidence))
            elif y + h < cy <= y + h + band:
                caption.append(ln)
        fields = extract_fields(inside, w, h)
        printed = _parse_caption(caption)
        fields.update({k: v for k, v in printed.items() if v not in ("", None)})
        fields["printed"] = sum(printed[k] not in ("", None) for k in ("check_no", "date", "amount"))
        if fields["printed"]:
            fields["value_conf"] = round(_payee_conf(fields, inside, w, h), 3)
        cap_text = " | ".join(ln.text for ln in sorted(caption, key=lambda l: l.bbox[0]))
        fields["raw_text"] = (fields["raw_text"] + f"\n[printed] {cap_text}").strip()
        out.append((fields, (x, y, w, int(h + band))))
    return out


def _caption_band(boxes, i):
    """Height of the gap below box i: up to the next box in its column, else half a check."""
    x, y, w, h = boxes[i]
    below = [b[1] for b in boxes if b[1] > y + h / 2 and abs(b[0] - x) < w / 2]
    return min(below) - (y + h) if below else 0.5 * h


def _parse_caption(caption):
    text = " ".join(ln.text for ln in sorted(caption, key=lambda l: (l.bbox[1] // 20, l.bbox[0])))
    no = CAPTION_NO_RE.search(text)
    date = DATE_RE.search(text)
    amt = CAPTION_AMT_RE.search(text)
    return {
        "check_no": no.group(1) if no else "",
        "date": date.group(0) if date else "",
        "amount": float(amt.group(1).replace(",", "")) if amt else None,
    }


def _payee_conf(fields, inside, w, h):
    """OCR confidence of the payee: re-derive it from the in-check extraction."""
    nl = [_norm(ln, w, h) for ln in inside]
    src, used = {}, set()
    _find_amount(nl, used, src)
    if not (_payee_from_amount_row(nl, used, src) or _value_after_label(nl, PAYEE_LABEL_RE, used, src, "payee")):
        return 0.0
    return src.get("payee", 0.0)


# ── Confidence ────────────────────────────────────────────────────────────────

def grade(fields):
    """HIGH / MEDIUM / LOW from field completeness + OCR confidence of amount/payee."""
    if not fields["raw_text"]:
        return "LOW"
    printed = fields.get("printed", 0)  # caption fields found (statement pages)
    if printed == 3:  # check no./date/amount are printed text; only the payee is uncertain
        if fields["payee"] and fields["value_conf"] >= 0.85:
            return "HIGH"
        return "MEDIUM" if fields["payee"] and fields["value_conf"] >= 0.6 else "LOW"
    if printed:
        return "LOW"  # a printed line was there but couldn't be fully read
    found = sum(bool(fields[k]) for k in ("check_no", "date", "payee", "purpose"))
    has_amount = fields["amount"] is not None
    if has_amount and fields["payee"] and fields["date"] and fields["value_conf"] >= 0.85:
        return "HIGH"
    if has_amount and found >= 2 and fields["value_conf"] >= 0.6:
        return "MEDIUM"
    return "LOW"


def apply_payee_match(fields, matcher):
    """Phase 2: swap the OCR payee for the client's vendor (see payee_match.py).

    Keeps the OCR text in payee_ocr. An alias hit is a confirmed reading, so the
    payee counts as fully confident; a fuzzy match leaves grading unchanged.
    """
    fields["payee_ocr"] = fields.get("payee", "")
    m = matcher.match(fields["payee_ocr"])
    fields["payee_match"] = m.kind
    if m.vendor:
        fields["payee"] = m.vendor
    if m.kind == "alias":  # value_conf is the payee's alone on statement pages
        fields["value_conf"] = 1.0 if fields.get("printed") else fields.get("amount_conf", fields["value_conf"])
    return m


def build_row(check, fields, ocr_failed=False):
    conf = "LOW" if ocr_failed else grade(fields)
    row = {
        "Source": check.source,
        "Page": check.page,
        "Check #": check.index,
        "Check No.": fields.get("check_no", ""),
        "Date": fields.get("date", ""),
        "Amount": fields.get("amount"),
        "Payee": "OCR FAILED" if ocr_failed else fields.get("payee", ""),
    }
    if "payee_ocr" in fields:  # payee matching ran (a client was picked)
        row["Payee (OCR)"] = fields["payee_ocr"]
    row.update({
        "Purpose": fields.get("purpose", ""),
        "Confidence": conf,
        "Flag": ocr_failed or conf == "LOW" or fields.get("payee_match") == "none",
    })
    return row


# ── Export ────────────────────────────────────────────────────────────────────

EXPORT_COLS = ["Source", "Page", "Check #", "Check No.", "Date", "Amount",
               "Payee", "Payee (OCR)", "Purpose", "Confidence", "Flag"]
USD_FORMAT = '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)'  # same as stock_processor
CONF_FILLS = {"HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "LOW": "FFC7CE"}


def to_excel(df):
    """Excel bytes: bold frozen header, filters, USD amounts, colour-coded confidence."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    df = df[[c for c in EXPORT_COLS if c in df.columns]]
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Checks")
        ws = writer.sheets["Checks"]
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        cols = list(df.columns)
        for idx, name in enumerate(cols, start=1):
            letter = get_column_letter(idx)
            width = max([len(str(name))] + [len(str(v)) for v in df[name].tolist()]) + 2
            ws.column_dimensions[letter].width = min(max(width, 8), 45)
        if "Amount" in cols:
            col = cols.index("Amount") + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).number_format = USD_FORMAT
        if "Confidence" in cols:
            col = cols.index("Confidence") + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, col)
                if cell.value in CONF_FILLS:
                    cell.fill = PatternFill("solid", fgColor=CONF_FILLS[cell.value])
    out.seek(0)
    return out.getvalue()
