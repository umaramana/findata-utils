"""
Synthetic tests for check_extractor (no real check images, no Surya needed).
Run: python bookkeeping/test_check_extractor.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import check_extractor as ce  # noqa: E402

W, H = 1200, 500  # ~2.4:1 check


def L(text, x0, y0, x1, y1, conf=0.95):
    return ce.OcrLine(text, (x0 * W, y0 * H, x1 * W, y1 * H), conf)


CASES = {
    # printed labels and handwriting detected as separate lines
    "separate_lines": [
        L("1042", 0.88, 0.05, 0.95, 0.10),
        L("DATE", 0.62, 0.17, 0.67, 0.21), L("3/15/2024", 0.69, 0.15, 0.85, 0.22),
        L("PAY TO THE ORDER OF", 0.04, 0.33, 0.25, 0.38), L("Acme Plumbing LLC", 0.28, 0.31, 0.60, 0.39),
        L("$", 0.78, 0.33, 0.80, 0.38), L("1,250.00", 0.81, 0.32, 0.95, 0.39),
        L("One thousand two hundred fifty and 00/100", 0.04, 0.46, 0.70, 0.52),
        L("DOLLARS", 0.86, 0.47, 0.95, 0.52),
        L("MEMO", 0.04, 0.75, 0.10, 0.80), L("Invoice 5531", 0.12, 0.73, 0.40, 0.81),
        L("⑆021000021⑆ 123456789⑈ 1042", 0.10, 0.88, 0.70, 0.95),
    ],
    # labels and values merged on one OCR line
    "merged_lines": [
        L("No. 2207", 0.85, 0.05, 0.96, 0.10),
        L("Date Jan 5, 2025", 0.60, 0.15, 0.90, 0.22),
        L("Pay to the order of John Smith $ 85.50", 0.04, 0.32, 0.95, 0.39),
        L("Eighty five and 50/100 dollars", 0.04, 0.46, 0.70, 0.52),
        L("For: lawn service", 0.04, 0.74, 0.40, 0.80),
    ],
    # handwriting noise: dashes in date, no $ sign, no memo, low conf
    "sparse_low_conf": [
        L("12-01-23", 0.68, 0.15, 0.85, 0.22, 0.55),
        L("PAY TO THE ORDER OF", 0.04, 0.33, 0.25, 0.38, 0.9),
        L("City Water Dept", 0.28, 0.31, 0.60, 0.39, 0.5),
        L("72.18", 0.82, 0.32, 0.95, 0.39, 0.6),
    ],
    # real Surya output (rendered check): label split over two lines, amount box "$ 1,250.00" one line
    "split_label_surya": [
        L("1042", .884, .052, .927, .080, .99), L("3/15/2024", .687, .155, .822, .202, .98),
        L("DATE", .630, .166, .687, .205, .99), L("PAY TO THE", .024, .348, .115, .378, .99),
        L("$ 1,250.00", .787, .350, .927, .405, .92), L("Acme Plumbing LLC", .127, .352, .397, .402, .97),
        L("ORDER OF", .024, .391, .105, .419, .97),
        L("One thousand two hundred fifty and 00/100", .023, .522, .582, .573, .99),
        L("DOLLARS", .891, .544, .965, .572, .67), L("MEMO Invoice 5531", .023, .731, .25, .781, .92),
    ],
    # split label with the handwriting aligned to the second label line
    "split_label_low_payee": [
        L("PAY TO", .02, .33, .09, .37), L("THE ORDER OF", .02, .38, .14, .42),
        L("Dr. Patel", .16, .37, .40, .43), L("$410", .80, .35, .95, .41), L("10/2/25", .70, .15, .85, .21),
    ],
    "empty": [],
}

EXPECT = {
    "separate_lines": dict(check_no="1042", date="3/15/2024", amount=1250.0,
                           payee="Acme Plumbing LLC", purpose="Invoice 5531", grade="HIGH"),
    "merged_lines": dict(check_no="2207", date="Jan 5, 2025", amount=85.5,
                         payee="John Smith", purpose="lawn service", grade="HIGH"),
    "sparse_low_conf": dict(check_no="", date="12-01-23", amount=72.18,
                            payee="City Water Dept", purpose="", grade="LOW"),
    "split_label_surya": dict(check_no="1042", date="3/15/2024", amount=1250.0,
                              payee="Acme Plumbing LLC", purpose="Invoice 5531", grade="HIGH"),
    "split_label_low_payee": dict(check_no="", date="10/2/25", amount=410.0,
                                  payee="Dr. Patel", purpose="", grade="HIGH"),
    "empty": dict(check_no="", date="", amount=None, payee="", purpose="", grade="LOW"),
}


def test_fields():
    fails = 0
    for name, lines in CASES.items():
        f = ce.extract_fields(lines, W, H)
        got = {k: f[k] for k in ("check_no", "date", "amount", "payee", "purpose")}
        got["grade"] = ce.grade(f)
        for k, v in EXPECT[name].items():
            if got[k] != v:
                fails += 1
                print(f"FAIL {name}.{k}: expected {v!r}, got {got[k]!r}")
    return fails


# Statement page: positions (% of a 2550x3299 page) and text shapes taken from the
# masked diag of a real 17-check page (06 Oct); names/words are invented stand-ins.
PW, PH = 2550, 3299
GRID_X, GRID_Y = (71, 881, 1691), (744, 1155, 1566)
STATEMENT_LINES = [  # x0, y0, x1, y1 (%), conf, text
    # row 1 — 961 / 962 / 975
    (28.5, 23.8, 30.1, 24.3, .98, "961"), (60.0, 23.8, 61.6, 24.3, .98, "962"), (91.7, 23.7, 93.4, 24.3, .98, "975"),
    (4.2, 23.6, 10.8, 25.2, .99, "JO'S FLOWER SHOP 7607 SC HWY 76"), (50.5, 24.6, 56.4, 25.6, .98, "01/13/26"),
    (4.1, 26.2, 18.7, 27.2, .95, "PAY TO THE ACME TILE"), (25.0, 26.3, 31.4, 27.2, .88, "$ 9,538.84"),
    (35.9, 26.2, 52.3, 27.2, .95, "PAY TO THE ACME STONE"), (56.4, 26.3, 62.3, 27.2, .94, "$ 709.28"),
    (67.7, 26.3, 73.7, 27.1, .96, "PAY TO THE DUKE"), (88.4, 26.4, 93.8, 27.2, .91, "$ 80568"),
    (3.1, 27.2, 31.3, 28.3, .94, "TWO THOUSAND FIVE HUNDRED THIRTY EIGHT AND 84/100 DOLLARS"),
    (66.8, 27.3, 88.1, 28.2, .93, "EIGHT HUNDRED FIVE AND 68/100"), (88.7, 27.7, 94.3, 28.3, .67, "_DOLLARS"),
    (4.2, 30.3, 12.1, 31.1, .99, "FOR A-629087"), (69.5, 30.1, 75.6, 30.9, .97, "FOR 1234567"),
    (4.3, 31.4, 21.6, 32.1, .88, "⑆053201607⑆ 123456789⑈0961"),
    (2.7, 33.4, 10.1, 34.3, .94, "Check# 961"), (15.1, 33.5, 22.0, 34.3, .99, "01/12/2026"), (24.5, 33.4, 30.2, 34.3, .98, "$2538.84"),
    (34.5, 33.4, 42.1, 34.3, .94, "Check# 962"), (46.8, 33.5, 53.7, 34.3, .99, "01/13/2026"), (56.2, 33.4, 61.4, 34.3, .97, "$709.28"),
    (66.2, 33.4, 73.8, 34.3, .95, "Check# 975"), (78.5, 33.4, 85.5, 34.3, .99, "01/05/2026"), (88.1, 33.4, 93.1, 34.3, .98, "$805.68"),
    # row 2 — 976 / 977 / 978
    (28.5, 36.3, 30.2, 36.9, .96, "976"),
    (4.2, 38.8, 15.1, 39.6, .92, "ORDER OF GREEN LAWN"), (24.8, 38.7, 31.5, 39.7, .97, "$6,193.27"),
    (36.1, 38.8, 51.0, 39.8, .93, "PAY TO THE PIEDMONT PLUMBER"), (56.4, 38.8, 62.4, 39.6, .91, "$ 2,097.36"),
    (67.8, 38.9, 82.3, 39.6, .62, "ORDER OF SPARTAN ROOFING"), (88.0, 38.8, 94.0, 39.7, .91, "$ 2511.07"),
    (2.7, 45.9, 10.3, 46.7, .95, "Check# 976"), (15.0, 45.9, 22.0, 46.8, .99, "01/05/2026"), (24.5, 45.9, 30.2, 46.8, .97, "$6193.27"),
    (34.5, 45.9, 42.0, 46.7, .94, "Check# 977"), (46.8, 45.9, 53.7, 46.8, .99, "01/07/2026"), (56.2, 45.9, 62.1, 46.8, .97, "$2097.36"),
    (66.2, 45.9, 73.8, 46.7, .94, "Check# 978"), (78.5, 45.9, 85.5, 46.8, .99, "01/14/2026"),  # amount unreadable
    # row 3 — 979 (label + name split over two OCR lines) / 980 (no printed line at all)
    (3.9, 51.3, 10.4, 52.1, .94, "PAY TO THE ACE CO"), (24.2, 51.2, 30.2, 52.2, .93, "1$7,645.49"),
    (35.9, 51.1, 46.7, 52.0, .97, "PAY TO"), (47.0, 51.3, 54.3, 52.0, .96, "Upstate Electric group"),
    (56.7, 51.5, 62.1, 52.1, .81, "$ 2,753,48"),
    (2.7, 58.4, 10.2, 59.2, .95, "Check# 979"), (15.0, 58.4, 22.1, 59.2, .99, "01/12/2026"), (24.5, 58.4, 30.2, 59.2, .98, "$7645.49"),
]
STATEMENT_EXPECT = [  # check_no, date, amount, payee, grade
    ("961", "01/12/2026", 2538.84, "ACME TILE", "HIGH"),   # printed $2538.84 beats handwritten 9,538.84
    ("962", "01/13/2026", 709.28, "ACME STONE", "HIGH"),
    ("975", "01/05/2026", 805.68, "DUKE", "HIGH"),
    ("976", "01/05/2026", 6193.27, "GREEN LAWN", "HIGH"),
    ("977", "01/07/2026", 2097.36, "PIEDMONT PLUMBER", "HIGH"),
    ("978", "01/14/2026", 2511.07, "SPARTAN ROOFING", "LOW"),  # printed amount missing -> LOW
    ("979", "01/12/2026", 7645.49, "ACE CO", "HIGH"),
    ("", "", 2753.48, "Upstate Electric group", "LOW"),        # no printed line: handwriting only
]


def test_statement_page():
    lines = [ce.OcrLine(t, (x0 * PW / 100, y0 * PH / 100, x1 * PW / 100, y1 * PH / 100), c)
             for x0, y0, x1, y1, c, t in STATEMENT_LINES]
    boxes = [(x, y, 750, 347) for y in GRID_Y for x in GRID_X][:8]
    results = ce.extract_page(lines, boxes)
    fails = 0
    for (f, view), exp in zip(results, STATEMENT_EXPECT):
        got = (f["check_no"], f["date"], f["amount"], f["payee"], ce.grade(f))
        if got != exp:
            fails += 1
            print(f"FAIL statement {exp[0] or '(no caption)'}: expected {exp}, got {got}")
    if results[0][1][3] <= 347:
        fails += 1
        print("FAIL statement: review image should include the printed line")
    return fails


def test_segmentation():
    import numpy as np
    from PIL import Image, ImageDraw
    fails = 0
    # letter page at 100 dpi with 3 checks (6" x 2.75") stacked
    page = Image.new("RGB", (850, 1100), "white")
    d = ImageDraw.Draw(page)
    for top in (60, 400, 740):
        d.rectangle((125, top, 725, top + 275), outline="black", width=3)
        d.rectangle((600, top + 90, 710, top + 125), outline="black", width=2)  # amount box (nested)
        d.text((140, top + 100), "PAY TO THE ORDER OF", fill="black")
    crops = ce.segment_checks(page)
    if len(crops) != 3:
        fails += 1
        print(f"FAIL segmentation: expected 3 checks, got {len(crops)}")
    elif not all(1.8 < c.width / c.height < 3.4 for c in crops):
        fails += 1
        print("FAIL segmentation: crop aspect ratios", [c.size for c in crops])
    # statement grid, geometry from a real 17-check page (2550x3299 px, diag 06 Oct)
    xs, ys = (71, 881, 1691), (744, 1155, 1566, 1977, 2388, 2799)
    for n_checks, name in ((17, "17-up grid"), (18, "18-up grid")):
        page = Image.new("RGB", (2550, 3299), "white")
        d = ImageDraw.Draw(page)
        d.rectangle((2095, 438, 2419, 513), outline="black", width=3)  # header bar, not a check
        d.rectangle((60, 80, 460, 260), outline="black", width=3)      # check-shaped logo, ~0.3x a check
        cells = [(x, y) for y in ys for x in xs][:n_checks]
        for x, y in cells:
            d.rectangle((x, y + (x % 7), x + 750, y + 347 + (x % 7)), outline="black", width=3)  # slight skew
        crops = ce.segment_checks(page)
        if len(crops) != n_checks:
            fails += 1
            print(f"FAIL {name}: expected {n_checks}, got {len(crops)}")
    # row-major order: shade each cell by its index, read the shades back from the crops
    page = Image.new("RGB", (2550, 3299), "white")
    d = ImageDraw.Draw(page)
    for i, (x, y) in enumerate([(x, y) for y in ys for x in xs][:17]):
        d.rectangle((x, y + (x % 7), x + 750, y + 347 + (x % 7)), outline="black", width=3)
        d.rectangle((x + 300, y + 150, x + 450, y + 200), fill=(10 * i, 10 * i, 10 * i))
    got = [c.getpixel((c.width // 2, c.height // 2))[0] // 10 for c in ce.segment_checks(page)]
    if got != list(range(17)):
        fails += 1
        print(f"FAIL grid order: {got}")
    # 24-up (4x6) and 1-up pages
    for cols, rows_, name in ((4, 6, "24-up"), (1, 1, "1-up")):
        page = Image.new("RGB", (2550, 3299), "white")
        d = ImageDraw.Draw(page)
        cw, ch = (560, 250) if cols == 4 else (2200, 1000)
        for r in range(rows_):
            for c in range(cols):
                x, y = 60 + c * (cw + 50), 300 + r * (ch + 150)
                d.rectangle((x, y, x + cw, y + ch), outline="black", width=3)
        n = len(ce.segment_checks(page))
        if n != cols * rows_:
            fails += 1
            print(f"FAIL {name}: expected {cols * rows_}, got {n}")
    # blank page -> fallback to whole page
    blank = Image.new("RGB", (850, 1100), "white")
    if ce.segment_checks(blank)[0].size != blank.size:
        fails += 1
        print("FAIL segmentation fallback")
    # photo of a single check (no border): fallback keeps the whole image
    single = Image.fromarray(np.full((500, 1200, 3), 240, np.uint8))
    if len(ce.segment_checks(single)) != 1:
        fails += 1
        print("FAIL single check")
    return fails


def test_excel():
    import pandas as pd
    from openpyxl import load_workbook
    import io
    df = pd.DataFrame([
        ce.build_row(ce.CheckImage("a.pdf", 1, 1, None), ce.extract_fields(CASES["separate_lines"], W, H)),
        ce.build_row(ce.CheckImage("a.pdf", 1, 2, None), ce.extract_fields([], W, H), ocr_failed=True),
    ])
    ws = load_workbook(io.BytesIO(ce.to_excel(df)))["Checks"]
    fails = 0
    if ws.freeze_panes != "A2" or not ws.auto_filter.ref:
        fails += 1
        print("FAIL excel header/filter")
    if ws["F2"].value != 1250.0 or "$" not in ws["F2"].number_format:
        fails += 1
        print("FAIL excel amount", ws["F2"].value, ws["F2"].number_format)
    if ws["I3"].value != "LOW" or ws["I3"].fill.fgColor.rgb[-6:] != "FFC7CE" or ws["G3"].value != "OCR FAILED":
        fails += 1
        print("FAIL excel OCR-failed row")
    return fails


if __name__ == "__main__":
    total = test_fields() + test_statement_page() + test_segmentation() + test_excel()
    print("ALL PASS" if total == 0 else f"{total} FAILURE(S)")
    sys.exit(1 if total else 0)
