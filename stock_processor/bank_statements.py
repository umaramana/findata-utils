"""
Bank Statements — pure logic for the page (no Streamlit): pick parser by format,
reconcile, fill check payees, build the Summary -> Master -> per-month workbook.
"""
import io
import os
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_SCRIPTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                        'bankdetails_dataextraction', 'scripts')
if _SCRIPTS not in sys.path:
    sys.path.append(_SCRIPTS)

import extract_regions_txns as regions          # noqa: E402
import extract_regions_cc_txns as regions_cc    # noqa: E402
from check_register import fill_checks          # noqa: E402

import bank_formats_legacy as legacy         # noqa: E402


def _regions(mod):
    def parse(pages):
        res = mod.parse_statement(pages)
        return {'transactions': res['transactions'], 'unparsed': res['unparsed'],
                'totals': mod.reconcile(res)}
    return parse


# format label -> parse(pages) -> {transactions, unparsed, totals}
FORMATS = {
    'Regions checking / savings': _regions(regions),
    'Regions credit card': _regions(regions_cc),
    'Chase business checking': legacy.parse_chase_checking,
    'Chase credit card': legacy.parse_chase_cc,
    'Citi checking': legacy.parse_citi,
}

_GREEN = PatternFill('solid', fgColor='C6EFCE')
_RED = PatternFill('solid', fgColor='FFC7CE')
_HDR_FILL = PatternFill('solid', fgColor='1F4E79')
_HDR_FONT = Font(color='FFFFFF', bold=True)
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_COLS = ['Date', 'Description', 'Amount', 'Section', 'Status', 'Filename']


def parse_file(fmt, pages):
    """Parse one statement. Returns {transactions, unparsed, totals: {name: (printed, extracted, ok)}}."""
    return FORMATS[fmt](pages)


def apply_register(txns, register=None):
    return fill_checks(txns, register)


def _month_key(d):
    return d.strftime('%Y-%m')


def _write_summary(ws, by_month, statements):
    """Old extractors' layout: Month | Transactions | one column per section | Net, then TOTAL.
    Below it, one block checking each statement's printed totals (green / red)."""
    sections = []
    for txns in by_month.values():
        for t in txns:
            if t['section'] not in sections:
                sections.append(t['section'])
    headers = ['Month', 'Transactions'] + sections + ['Net']
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h)
        c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _CENTER
    ws.row_dimensions[1].height = 36

    def put_row(r, label, txns, bold=False):
        ws.cell(r, 1, label)
        ws.cell(r, 2, len(txns))
        net = 0.0
        for col, sec in enumerate(sections, 3):
            val = round(sum(t['amount'] for t in txns if t['section'] == sec), 2)
            ws.cell(r, col, val).number_format = '#,##0.00'
            net += val
        ws.cell(r, len(headers), round(net, 2)).number_format = '#,##0.00'
        if bold:
            for col in range(1, len(headers) + 1):
                ws.cell(r, col).font = Font(bold=True)

    months = sorted(by_month)
    for r, m in enumerate(months, 2):
        put_row(r, m, by_month[m])
    put_row(len(months) + 2, 'TOTAL', [t for m in months for t in by_month[m]], bold=True)

    r = len(months) + 5
    for col, h in enumerate(['Statement', 'Check', 'Printed', 'Extracted', 'Gap', 'Result'], 1):
        c = ws.cell(r, col, h)
        c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _CENTER
    for s in statements:
        for name, (printed, extracted, ok) in s['totals'].items():
            r += 1
            gap = None if printed is None else round(extracted - printed, 2)
            for col, v in enumerate([s['name'], name, printed, extracted, gap,
                                     'no printed total' if ok is None else ('OK' if ok else 'MISMATCH')], 1):
                ws.cell(r, col, v)
            if ok is not None:
                ws.cell(r, 6).fill = _GREEN if ok else _RED
        if s.get('unparsed'):
            r += 1
            for col, v in enumerate([s['name'], 'Unparsed lines', None, len(s['unparsed']), None, 'REVIEW'], 1):
                ws.cell(r, col, v)
            ws.cell(r, 6).fill = _RED
    for row in ws.iter_rows(min_row=1, min_col=3, max_col=5):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00'
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 24
    for i in range(len(sections) + 1):
        ws.column_dimensions[get_column_letter(3 + i)].width = 22


def build_workbook(statements):
    """
    statements: [{'name', 'transactions', 'totals'}] (transactions already status-filled).
    Returns BytesIO: Summary, Master (via the shared collator), then one tab per month.
    """
    import excel_collator as C

    by_month = defaultdict(list)
    for s in statements:
        for t in s['transactions']:
            by_month[_month_key(t['date'])].append({**t, 'file': s['name']})

    wb = Workbook()
    wb.remove(wb.active)
    for month in sorted(by_month):
        ws = wb.create_sheet(month)
        ws.append(_COLS)
        for t in sorted(by_month[month], key=lambda x: x['date']):
            ws.append([t['date'].strftime('%m/%d/%Y'), t['description'], f"{t['amount']:.2f}",
                       t['section'], t.get('status', 'OK'), t['file']])
    if not wb.sheetnames:
        wb.create_sheet('Empty').append(_COLS)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    df = C._collate_sheets(buf, wb.sheetnames)
    types = {'Month': 'text', 'Date': 'date', 'Description': 'text',
             'Amount': 'currency', 'Section': 'text', 'Status': 'text', 'Filename': 'text'}
    buf.seek(0)
    out = C._generate_output(buf, df, types)

    wb2 = load_workbook(out)
    ws = wb2.create_sheet('Summary')
    _write_summary(ws, by_month, statements)
    wb2._sheets.remove(ws)
    wb2._sheets.insert(0, ws)
    final = io.BytesIO()
    wb2.save(final)
    final.seek(0)
    return final
