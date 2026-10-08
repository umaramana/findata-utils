"""Tests for the shared Excel collator. Run: python3 test_excel_collator.py"""
import io
import unittest

from openpyxl import Workbook, load_workbook

import excel_collator as C


def _book():
    wb = Workbook()
    a = wb.active
    a.title = 'Jan'
    a.append(['Date', 'Description', 'Amount'])
    a.append(['01/05/2026', 'ACME Store', '100.50'])
    a.append(['01/06/2026', 'KLMN Supply', '-20.00'])
    b = wb.create_sheet('Feb')
    b.append(['Date', 'Description', 'Amount'])
    b.append(['02/01/2026', 'ACME Store', '10.00'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class Collator(unittest.TestCase):
    def test_collate_adds_month_and_skips_repeat_headers(self):
        df = C._collate_sheets(_book(), ['Jan', 'Feb'])
        self.assertEqual(list(df.columns), ['Month', 'Date', 'Description', 'Amount'])
        self.assertEqual(list(df['Month']), ['Jan', 'Jan', 'Feb'])

    def test_output_master_first_with_recon(self):
        buf = _book()
        df = C._collate_sheets(buf, ['Jan', 'Feb'])
        buf.seek(0)
        types = {'Month': 'text', 'Date': 'date', 'Description': 'text', 'Amount': 'currency'}
        out = C._generate_output(buf, df, types)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames[0], 'Master')
        rows = [r for r in wb['Master'].iter_rows(values_only=True)]
        grand = [r for r in rows if r[0] == 'Grand Total'][0]
        self.assertEqual(grand[1:3], (3, 90.5))


if __name__ == '__main__':
    unittest.main()
