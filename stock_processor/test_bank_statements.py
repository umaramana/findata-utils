"""Tests for bank_statements (invented names). Run: python3 test_bank_statements.py"""
import unittest

from openpyxl import load_workbook

import bank_statements as B

PAGE = """ACME BUSINESS CHECKING January 14, 2026 through February 10, 2026
DEPOSITS & CREDITS
01/15 Deposit ACME Store 7 1,000.00
Total Deposits & Credits $1,000.00
CHECKS
02/02 911 40.00 02/03 912 60.00
Total Checks $100.00
"""


class Page(unittest.TestCase):
    def test_end_to_end(self):
        p = B.parse_file('Regions checking / savings', [PAGE])
        B.apply_register(p['transactions'], {'911': ('KLMN Supply', 40.0)})
        statuses = {t['description']: t['status'] for t in p['transactions']}
        self.assertEqual(statuses['KLMN Supply'], 'payee filled')
        self.assertEqual(statuses['Check #912'], 'check not in register')
        wb = load_workbook(B.build_workbook([{'name': 's.pdf', **p}]))
        self.assertEqual(wb.sheetnames, ['Summary', 'Master', '2026-01', '2026-02'])
        summary = list(wb['Summary'].iter_rows(values_only=True))
        self.assertEqual(summary[0][:2], ('Month', 'Transactions'))
        self.assertIn('Net', summary[0])
        self.assertEqual([r[0] for r in summary[1:4]], ['2026-01', '2026-02', 'TOTAL'])
        self.assertEqual(summary[3][summary[0].index('Net')], 900.0)
        self.assertTrue(all(r[5] == 'OK' for r in summary if r[1] in ('Checks', 'Deposits & Credits')))
        master = list(wb['Master'].iter_rows(values_only=True))
        self.assertEqual(master[0][:6], ('Month', 'Date', 'Description', 'Amount', 'Section', 'Status'))

    def test_gap_marked_red(self):
        p = B.parse_file('Regions checking / savings', [PAGE.replace('40.00', '41.00')])
        wb = load_workbook(B.build_workbook([{'name': 's.pdf', **p}]))
        res = {r[1]: r[5] for r in wb['Summary'].iter_rows(values_only=True)}
        self.assertEqual(res['Checks'], 'MISMATCH')


CHASE_CC = """Opening/Closing Date 12/20/25 - 01/19/26
ACCOUNT ACTIVITY
12/22 ACME STORE 10.00
01/05 KLMN SUPPLY 5.50
TRANSACTIONS THIS CYCLE $15.50
"""
CITI = """\u00a9 2026
CHECKING ACTIVITY
01/05 Deposit ACME 100.00 600.00
Total Subtracted/Added 0.00 100.00
"""
CHASE_CK = """ACME BUSINESS CHECKING December 20, 2025 through January 19, 2026
DEPOSITS AND ADDITIONS
12/22 Deposit 10.00
01/05 Deposit 5.00
Total Deposits and Additions $15.00
"""


class Legacy(unittest.TestCase):
    def test_chase_cc_year_rollover_and_total(self):
        p = B.parse_file('Chase credit card', [CHASE_CC])
        self.assertEqual([str(t['date']) for t in p['transactions']], ['2025-12-22', '2026-01-05'])
        self.assertTrue(p['totals']['Transactions this cycle'][2])

    def test_chase_checking_rollover(self):
        p = B.parse_file('Chase business checking', [CHASE_CK])
        self.assertEqual([str(t['date']) for t in p['transactions']], ['2025-12-22', '2026-01-05'])

    def test_citi_runs(self):
        p = B.parse_file('Citi checking', [CITI])
        self.assertEqual(len(p['transactions']) + len(p['unparsed']), 1)

    def test_preset_status_survives_register_fill(self):
        t = [{'description': 'x', 'amount': 1.0, 'status': 'date unreadable'}]
        B.apply_register(t, None)
        self.assertEqual(t[0]['status'], 'date unreadable')


if __name__ == '__main__':
    unittest.main()
